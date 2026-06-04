"""Export XLSX du résultat d'analyse (openpyxl).

Feuilles produites :
- ``Paramètres``      : contexte de l'export (date, couches, champs, filtres) ;
- ``Liste poteaux``   : détail des poteaux sortis (table structurée) ;
- ``Câbles associés`` : détail des câbles rattachés aux poteaux sortis ;
- ``Synthèse``        : tableaux récapitulatifs calculés côté Python ;
- ``TCD``             : table à plat prête à pivoter.

L'export ne dépend pas de QGIS : il peut être généré et testé hors QGIS.

``openpyxl`` est importé de façon paresseuse : son absence n'empêche pas le
chargement du plugin dans QGIS, seul l'export lève alors une erreur explicite.
"""

from datetime import datetime

from .columns import (
    POLE_COLUMNS, CABLE_COLUMNS, GC_COLUMNS, CABLE_GC_COLUMNS,
)
from .synthese import (
    build_syntheses, build_tcd_table, TCD_COLUMNS, GC_TCD_COLUMNS,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    _OPENPYXL_IMPORT_ERROR = None
except ImportError as exc:  # openpyxl absent de l'environnement QGIS
    _OPENPYXL_IMPORT_ERROR = exc


def _require_openpyxl():
    if _OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError(
            "Le module « openpyxl » est requis pour l'export XLSX mais n'est "
            "pas installé dans l'environnement Python de QGIS.\n\n"
            "Pour l'installer, ouvrez la Console Python de QGIS et exécutez :\n"
            "    import pip; pip.main(['install', 'openpyxl'])\n"
            "puis redémarrez QGIS."
        )


def _header_fill():
    return PatternFill("solid", fgColor="1F4E78")


def _header_font():
    return Font(bold=True, color="FFFFFF")


def _title_font():
    return Font(bold=True, size=12, color="1F4E78")


def _style_header(ws, row, ncols):
    fill = _header_fill()
    font = _header_font()
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws, ncols, max_width=60):
    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        width = 10
        for cell in ws[letter]:
            value = cell.value
            if value is not None:
                width = max(width, min(max_width, len(str(value)) + 2))
        ws.column_dimensions[letter].width = width


def _write_table_sheet(ws, columns, data, table_name):
    """Écrit une feuille avec en-tête figé, auto-filtre et table structurée."""
    labels = [label for _, label in columns]
    keys = [key for key, _ in columns]

    ws.append(labels)
    _style_header(ws, 1, len(labels))

    for row in data:
        ws.append([row.get(key, "") for key in keys])

    nrows = len(data) + 1
    ncols = len(labels)
    last_cell = "{}{}".format(get_column_letter(ncols), nrows)

    # Une table Excel structurée exige au moins une ligne de données.
    if data:
        table = Table(displayName=table_name, ref="A1:{}".format(last_cell))
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = "A1:{}".format(get_column_letter(ncols) + "1")

    ws.freeze_panes = "A2"
    _autosize(ws, ncols)


def _write_params_sheet(ws, params):
    ws["A1"] = "Paramètres de l'analyse CEM NCVL"
    ws["A1"].font = _title_font()
    ws.append([])
    ws.append(["Paramètre", "Valeur"])
    _style_header(ws, 3, 2)
    for label, value in params:
        if isinstance(value, (list, tuple, set)):
            value = " ; ".join(str(v) for v in value)
        ws.append([label, value])
    ws.freeze_panes = "A4"
    _autosize(ws, 2, max_width=80)


def _write_synthese_sheet(ws, syntheses):
    ws["A1"] = "Synthèses"
    ws["A1"].font = _title_font()
    current = 3
    for title, pairs in syntheses.items():
        ws.cell(row=current, column=1, value=title).font = Font(bold=True)
        current += 1
        ws.cell(row=current, column=1, value="Valeur")
        ws.cell(row=current, column=2, value="Nombre")
        _style_header(ws, current, 2)
        current += 1
        for label, nb in pairs:
            ws.cell(row=current, column=1, value=label)
            ws.cell(row=current, column=2, value=nb)
            current += 1
        current += 1  # ligne vide entre deux tableaux
    _autosize(ws, 2, max_width=60)


def _write_tcd_sheet(ws, tcd_rows, dim_columns, measure_key, measure_label,
                     table_name):
    labels = [label for _, label in dim_columns] + [measure_label]
    keys = [key for key, _ in dim_columns] + [measure_key]
    ws.append(labels)
    _style_header(ws, 1, len(labels))
    for row in tcd_rows:
        ws.append([row.get(key, "") for key in keys])
    ncols = len(labels)
    if tcd_rows:
        last = "{}{}".format(get_column_letter(ncols), len(tcd_rows) + 1)
        table = Table(displayName=table_name, ref="A1:{}".format(last))
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium9", showRowStripes=True,
        )
        ws.add_table(table)
    else:
        ws.auto_filter.ref = "A1:{}1".format(get_column_letter(ncols))
    ws.freeze_panes = "A2"
    _autosize(ws, ncols)


def _assemble_workbook(rows, cable_detail, params, *, list_columns,
                       list_sheet, list_table, cable_columns, cable_sheet,
                       cable_table, tcd_dims, tcd_measure_key,
                       tcd_measure_label, tcd_table):
    """Assemble un classeur 5 feuilles, paramétré par jeu de colonnes.

    Mutualise la mise en forme entre l'export poteaux et l'export GC.
    """
    _require_openpyxl()
    wb = Workbook()

    ws_params = wb.active
    ws_params.title = "Paramètres"
    _write_params_sheet(ws_params, params)

    _write_table_sheet(
        wb.create_sheet(list_sheet), list_columns, rows, list_table)
    _write_table_sheet(
        wb.create_sheet(cable_sheet), cable_columns, cable_detail, cable_table)
    _write_synthese_sheet(
        wb.create_sheet("Synthèse"), build_syntheses(rows, cable_detail))
    _write_tcd_sheet(
        wb.create_sheet("TCD"),
        build_tcd_table(rows, tcd_dims, tcd_measure_key),
        tcd_dims, tcd_measure_key, tcd_measure_label, tcd_table)
    return wb


def build_workbook(rows, cable_detail, params):
    """Construit le classeur de l'analyse poteaux (sans l'enregistrer)."""
    return _assemble_workbook(
        rows, cable_detail, params,
        list_columns=POLE_COLUMNS, list_sheet="Liste poteaux",
        list_table="ListePoteaux",
        cable_columns=CABLE_COLUMNS, cable_sheet="Câbles associés",
        cable_table="CablesAssocies",
        tcd_dims=TCD_COLUMNS, tcd_measure_key="nb_poteaux",
        tcd_measure_label="Nb poteaux", tcd_table="TCD_Poteaux")


def build_gc_workbook(rows, cable_detail, params):
    """Construit le classeur de l'analyse GC souterrain (sans l'enregistrer)."""
    return _assemble_workbook(
        rows, cable_detail, params,
        list_columns=GC_COLUMNS, list_sheet="GC souterrains",
        list_table="ListeGC",
        cable_columns=CABLE_GC_COLUMNS, cable_sheet="Câbles associés",
        cable_table="CablesAssociesGC",
        tcd_dims=GC_TCD_COLUMNS, tcd_measure_key="nb_gc",
        tcd_measure_label="Nb GC", tcd_table="TCD_GC")


def export_xlsx(path, rows, cable_detail, params):
    """Génère le fichier XLSX (analyse poteaux) à ``path``."""
    build_workbook(rows, cable_detail, params).save(path)
    return path


def export_gc_xlsx(path, rows, cable_detail, params):
    """Génère le fichier XLSX (analyse GC souterrain) à ``path``."""
    build_gc_workbook(rows, cable_detail, params).save(path)
    return path


def build_params(context):
    """Met en forme le contexte d'export en liste (label, valeur) ordonnée.

    ``context`` est un dict libre fourni par l'interface ; les clés connues
    sont ordonnées, les autres ajoutées ensuite.
    """
    context = dict(context or {})
    ordered_labels = [
        ("date_export", "Date d'export"),
        ("couche_poteaux", "Couche poteaux"),
        ("couche_cables", "Couche câbles"),
        ("champ_id_poteau", "Champ ID poteau"),
        ("champ_etat_poteau", "Champ état poteau"),
        ("champ_statut_cable", "Champ statut câble"),
        ("champ_ref_cable", "Champ référence câble"),
        ("etats_poteau_retenus", "États poteau retenus"),
        ("statuts_tires", "Statuts câble « tirés »"),
        ("buffer_m", "Rayon buffer (m)"),
        ("crs_analyse", "CRS d'analyse"),
        ("poteaux_total", "Poteaux analysés"),
        ("poteaux_selectionnes", "Poteaux remplacés / implantés"),
        ("cables_total", "Câbles analysés"),
        ("poteaux_sans_cable_tire", "Poteaux sans câble tiré"),
    ]
    return _ordered_params(context, ordered_labels)


def build_gc_params(context):
    """Met en forme le contexte d'export GC en liste (label, valeur) ordonnée."""
    ordered_labels = [
        ("date_export", "Date d'export"),
        ("couche_gc", "Couche GC"),
        ("couche_cables", "Couche câbles"),
        ("champ_id_gc", "Champ ID GC"),
        ("champ_nom_gc", "Champ nom / code GC"),
        ("champ_suivi", "Champ suivi travaux GC"),
        ("etats_gc_retenus", "États GC retenus (travaux faits)"),
        ("territoires_retenus", "Territoires / plaques retenus"),
        ("statuts_tires", "Statuts câble « tirés »"),
        ("buffer_m", "Rayon buffer (m)"),
        ("crs_analyse", "CRS d'analyse"),
        ("gc_total", "GC analysés"),
        ("gc_selectionnes", "GC travaux faits"),
        ("cables_total", "Câbles analysés"),
        ("gc_sans_cable_tire", "GC sans câble tiré"),
    ]
    return _ordered_params(dict(context or {}), ordered_labels)


def _ordered_params(context, ordered_labels):
    """Ordonne les clés connues, puis ajoute les autres ; injecte la date."""
    if "date_export" not in context:
        context["date_export"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    params = []
    used = set()
    for key, label in ordered_labels:
        if key in context:
            params.append((label, context[key]))
            used.add(key)
    for key, value in context.items():
        if key not in used:
            params.append((key, value))
    return params
