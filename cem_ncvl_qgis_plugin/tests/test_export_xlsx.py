from openpyxl import load_workbook

from cem_ncvl_qgis_plugin.core.export_xlsx import (
    build_params, build_gc_params, export_xlsx, export_gc_xlsx,
)


def test_export_creates_all_sheets(tmp_path):
    rows = [{
        "id_poteau": "A", "commune": "Bordeaux", "departement": "33",
        "territoire": "T1", "etat_poteau": "plante", "nb_cables": 1,
        "etats_cables": "en_cours", "cables_associes": "c1",
        "cables_non_tires": "en_cours", "motif": "Aucun câble au statut tiré",
        "rayon_buffer_m": 5.0, "nb_cables_intersectes": 1,
        "mode_rattachement": "Spatial",
    }]
    cable_detail = [{
        "id_poteau": "A", "ref_cable": "c1", "statut_cable": "en_cours",
        "distance_m": 3.0, "commune": "Bordeaux", "departement": "33",
        "territoire": "T1",
    }]
    params = build_params({
        "couche_poteaux": "11_TRAVAUX_SUPPORTS.geom",
        "couche_cables": "0_cable_suivi.geom",
        "buffer_m": 5.0,
        "statuts_tires": ["Tiré", "Tirage fini"],
    })

    path = tmp_path / "out.xlsx"
    export_xlsx(str(path), rows, cable_detail, params)

    wb = load_workbook(str(path))
    assert set(wb.sheetnames) == {
        "Paramètres", "Liste poteaux", "Câbles associés", "Synthèse", "TCD"}

    ws = wb["Liste poteaux"]
    assert ws["A1"].value == "ID poteau"
    assert ws["A2"].value == "A"
    assert ws.freeze_panes == "A2"


def test_export_with_empty_rows(tmp_path):
    params = build_params({"buffer_m": 5.0})
    path = tmp_path / "empty.xlsx"
    export_xlsx(str(path), [], [], params)
    wb = load_workbook(str(path))
    # même vide, toutes les feuilles attendues existent
    assert "Liste poteaux" in wb.sheetnames
    assert "TCD" in wb.sheetnames


def test_gc_export_creates_all_sheets(tmp_path):
    rows = [{
        "id_gc": "GC-A", "nom_gc": "GC-A", "commune": "Bordeaux",
        "departement": "33", "territoire": "P1", "suivi_pilotage": "TRX fini",
        "longueur": "120", "nb_cables": 1, "etats_cables": "en_cours",
        "cables_associes": "c1", "cables_non_tires": "en_cours",
        "motif": "Aucun câble au statut tiré sur GC", "rayon_buffer_m": 1.0,
        "nb_cables_intersectes": 1, "mode_rattachement": "Spatial",
    }]
    cable_detail = [{
        "id_gc": "GC-A", "ref_cable": "c1", "statut_cable": "en_cours",
        "distance_m": 0.0, "commune": "Bordeaux", "departement": "33",
        "territoire": "P1",
    }]
    params = build_gc_params({
        "couche_gc": "0_artere_gc", "buffer_m": 1.0,
        "statuts_tires": ["Tiré", "Tirage fini"],
    })
    path = tmp_path / "gc.xlsx"
    export_gc_xlsx(str(path), rows, cable_detail, params)

    wb = load_workbook(str(path))
    assert set(wb.sheetnames) == {
        "Paramètres", "GC souterrains", "Câbles associés", "Synthèse", "TCD"}
    ws = wb["GC souterrains"]
    assert ws["A1"].value == "ID GC"
    assert ws["A2"].value == "GC-A"


def test_build_params_orders_known_keys_first():
    params = build_params({"buffer_m": 5.0, "couche_poteaux": "p"})
    labels = [label for label, _ in params]
    # 'date_export' est injecté et ordonné en tête
    assert labels[0] == "Date d'export"
    assert "Couche poteaux" in labels
    assert "Rayon buffer (m)" in labels
